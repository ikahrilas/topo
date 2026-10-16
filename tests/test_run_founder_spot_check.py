"""Seam 1: the founder spot-check loop, via `topo report` / `topo run` and `topo ingest-findings`.

`topo report` delivers a frozen spot-check copy of the pull's workbook; the founder marks rows
("Looks right?", Field, Note) and the Missing sheet; `topo ingest-findings` appends the allowed
columns to the write-once Review findings file; later pulls show "Last founder mark" by FAIN.
"""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook

from support.curated import override_row, write_overrides
from support.ipeds import DIRECTORY_COLUMNS, ENROLLMENT_COLUMNS, directory_row, enrollment_row, write_csv
from support.markdown import table
from support.usaspending import action, write_actions
from support.workbook import read_sheet, sheet_names, sheet_text

PULL = "2026-10-07"
REPO_ROOT = Path(__file__).resolve().parents[1]
SPOT_CHECK = f"spot_check_{PULL}.xlsx"
FINDINGS = Path("data") / "curated" / "review_findings.csv"
IHE = "PUBLIC/STATE CONTROLLED INSTITUTION OF HIGHER EDUCATION"


def one_year(fain, recipient, uei, state, amount=300000):
    return action(fain, 1, "2025-07-01", amount, ("2025-09-01", "2026-08-31"), action_type="A1",
                  recipient=recipient, uei=uei, state=state, business_types_description=IHE)


GRANTS = [
    one_year("P042A250001", "COLLEGE NUMBER 1", "MTUEI0000001", "MT"),
    one_year("P042A250002", "COLLEGE NUMBER 2", "MTUEI0000002", "MT"),
    one_year("P047A250003", "COLLEGE NUMBER 3", "MTUEI0000003", "MT"),
    one_year("P042A250004", "UNIVERSITY OF MINNESOTA DULUTH", "MNUEI0000004", "MN"),
    one_year("P042A250005", "OKLAHOMA COLLEGE", "OKUEI0000005", "OK"),
]


def setup(project):
    write_actions(project.raw_usaspending, PULL, GRANTS)
    folder = project.root / "data" / "raw" / "ipeds" / PULL
    write_csv(folder / "colleges_ipeds_directory.csv", DIRECTORY_COLUMNS, [
        directory_row("900001", "College Number 1", "MT", ueis=("MTUEI0000001",)),
        directory_row("900002", "College Number 2", "MT", ueis=("MTUEI0000002",)),
        directory_row("900003", "College Number 3", "MT", ueis=("MTUEI0000003",)),
        directory_row("174233", "University of Minnesota-Duluth", "MN", ueis=("MNUEI0000004",)),
    ])
    write_csv(folder / "colleges_ipeds_fall-enrollment-race_2024.csv", ENROLLMENT_COLUMNS,
              [enrollment_row("900001", 2024, 900)])


def by_fain(rows):
    return {r["FAIN"]: r for r in rows}


def mark(spot_check: Path, returned: Path, marks: dict[str, tuple], missing=(), extra_column=None) -> Path:
    """Fill a delivered spot-check copy as the founder would and save it as the returned workbook.

    marks: FAIN -> (Looks right?, Field, Note). missing: (State, Recipient, Program, Note) rows.
    extra_column: (header, value) written into a new column on every marked row (data the ingest must drop).
    """
    book = load_workbook(spot_check)
    sheet = book["Target list"]
    header = [c.value for c in sheet[1]]
    col = {name: header.index(name) + 1 for name in header}
    if extra_column:
        sheet.cell(row=1, column=len(header) + 1, value=extra_column[0])
    for row in range(2, sheet.max_row + 1):
        fain = sheet.cell(row=row, column=col["FAIN"]).value
        if fain in marks:
            looks, field, note = marks[fain]
            sheet.cell(row=row, column=col["Looks right?"], value=looks)
            sheet.cell(row=row, column=col["Field"], value=field)
            sheet.cell(row=row, column=col["Note"], value=note)
            if extra_column:
                sheet.cell(row=row, column=len(header) + 1, value=extra_column[1])
    for values in missing:
        book["Missing"].append(list(values))
    book.save(returned)
    return returned


def test_report_delivers_a_spot_check_copy_with_state_counts_mark_columns_and_a_names_warning(project):
    setup(project)

    project.topo("run", "--as-of", PULL)

    path = project.output / SPOT_CHECK
    assert sheet_names(path)[:2] == ["Spot check", "States"]
    assert {"Target list", "Ended", "Missing"} <= set(sheet_names(path))
    states = {r["State"]: r["Grants"] for r in read_sheet(path, "States")}
    assert states == {"MN": "1", "MT": "3", "OK": "1"}
    instructions = sheet_text(path, "Spot check")
    assert "Do not enter client or personal names" in instructions
    assert "ok / wrong / stale" in instructions
    assert "100 characters" in instructions
    rows = by_fain(read_sheet(path, "Target list"))
    assert set(rows) == {g["award_id_fain"] for g in GRANTS}
    assert {"Looks right?", "Field", "Note", "Last founder mark"} <= set(rows["P042A250001"])
    assert rows["P042A250001"]["Looks right?"] == ""
    assert list(read_sheet_header(path, "Missing")) == ["State", "Recipient", "Program", "Note"]
    assert (project.output / f"{SPOT_CHECK}.provenance.json").exists()


def read_sheet_header(path, sheet):
    return [c.value for c in next(load_workbook(path, read_only=True)[sheet].iter_rows(max_row=1))]


def test_the_spot_check_copy_locks_everything_but_the_mark_columns(project):
    setup(project)

    project.topo("run", "--as-of", PULL)

    sheet = load_workbook(project.output / SPOT_CHECK)["Target list"]
    header = [c.value for c in sheet[1]]
    assert sheet.protection.sheet
    editable = {header[c.column - 1] for c in sheet[2] if not c.protection.locked}
    assert editable == {"Looks right?", "Field", "Note"}


def test_a_rerun_never_overwrites_a_delivered_spot_check_copy(project):
    setup(project)
    project.topo("run", "--as-of", PULL)
    first = (project.output / SPOT_CHECK).read_bytes()

    project.topo("report")

    assert (project.output / SPOT_CHECK).read_bytes() == first
    assert (project.output / f"spot_check_{PULL}_v2.xlsx").exists()


def findings(project) -> list[dict[str, str]]:
    with (project.root / FINDINGS).open(newline="") as f:
        return list(csv.DictReader(f))


def returned(project, marks, **kw) -> Path:
    return mark(project.output / SPOT_CHECK, project.root / "returned.xlsx", marks, **kw)


def test_ingest_keeps_only_the_allowed_columns_with_reviewer_founder(project):
    setup(project)
    project.topo("run", "--as-of", PULL)
    path = returned(
        project,
        {
            "P042A250001": ("ok", None, None),
            "P042A250004": ("wrong", "Host institution", "Runs at the Twin Cities campus"),
            "P042A250005": ("stale", "Recipient", "Renamed in 2025"),
        },
        missing=[("MT", "Big Sky Tribal College", "Upward Bound", "Funded since 2022")],
        extra_column=("Contact", "Jane Q. Person, 555-0100"),
    )

    project.topo("ingest-findings", str(path), "--date", "2026-10-20")

    text = (project.root / FINDINGS).read_text()
    assert text.splitlines()[0] == "fain,state,recipient,program,mark,field,note,date,reviewer"
    assert "Jane Q. Person" not in text and "555-0100" not in text
    rows = findings(project)
    assert {r["reviewer"] for r in rows} == {"founder"}
    assert {r["date"] for r in rows} == {"2026-10-20"}
    by_key = {(r["fain"] or r["recipient"]): r for r in rows}
    assert set(by_key) == {"P042A250001", "P042A250004", "P042A250005", "Big Sky Tribal College"}
    assert by_key["P042A250004"] == {
        "fain": "P042A250004", "state": "MN", "recipient": "", "program": "", "mark": "wrong",
        "field": "Host institution", "note": "Runs at the Twin Cities campus", "date": "2026-10-20",
        "reviewer": "founder",
    }
    assert by_key["Big Sky Tribal College"] == {
        "fain": "", "state": "MT", "recipient": "Big Sky Tribal College", "program": "Upward Bound",
        "mark": "missing", "field": "", "note": "Funded since 2022", "date": "2026-10-20", "reviewer": "founder",
    }


def test_the_findings_file_is_append_only_and_a_rerun_adds_nothing(project):
    setup(project)
    (project.root / FINDINGS).parent.mkdir(parents=True)
    shutil.copy(REPO_ROOT / FINDINGS, project.root / FINDINGS)  # the shipped, header-only file
    project.topo("run", "--as-of", PULL)
    first = returned(project, {"P042A250001": ("ok", None, None)})
    project.topo("ingest-findings", str(first), "--date", "2026-10-20")
    before = (project.root / FINDINGS).read_bytes()

    project.topo("ingest-findings", str(first), "--date", "2026-10-20")
    assert (project.root / FINDINGS).read_bytes() == before

    second = mark(project.output / SPOT_CHECK, project.root / "second.xlsx",
                  {"P042A250001": ("stale", "Project end", "Extended a year")})
    project.topo("ingest-findings", str(second), "--date", "2026-11-02")
    after = (project.root / FINDINGS).read_bytes()
    assert after.startswith(before) and len(after) > len(before)
    assert [(r["mark"], r["date"]) for r in findings(project)] == [("ok", "2026-10-20"), ("stale", "2026-11-02")]


def test_ingest_refuses_a_mark_or_field_outside_the_lists_and_an_over_long_note(project):
    setup(project)
    project.topo("run", "--as-of", PULL)
    for marks in (
        {"P042A250001": ("maybe", None, None)},
        {"P042A250001": ("wrong", "Program officer", "x")},
        {"P042A250001": ("wrong", "Recipient", "x" * 101)},
    ):
        result = project.topo("ingest-findings", str(returned(project, marks)), "--date", "2026-10-20",
                              expect_exit=2)
        assert "Structural break" in result.output
    assert not (project.root / FINDINGS).exists() or findings(project) == []


def test_findings_never_change_the_overrides_file(project):
    setup(project)
    overrides = write_overrides(project.root, [override_row("P042A250001", "UNITID 900001")])
    before = overrides.read_bytes()
    project.topo("run", "--as-of", PULL)
    path = returned(project, {"P042A250004": ("wrong", "Host institution", "Runs at the Twin Cities campus")})

    project.topo("ingest-findings", str(path), "--date", "2026-10-20")
    project.topo("report")

    assert overrides.read_bytes() == before
    rows = by_fain(read_sheet(project.output / f"target_list_{PULL}_v2.xlsx", "Target list"))
    assert rows["P042A250004"]["Host institution"] == "University of Minnesota-Duluth"
    assert rows["P042A250004"]["Host confidence"] == "matched"
