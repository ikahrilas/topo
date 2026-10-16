"""Read a delivered target-list workbook back as plain rows."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import load_workbook


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def sheet_names(path: Path) -> list[str]:
    return load_workbook(path, read_only=True).sheetnames


def read_sheet(path: Path, sheet: str) -> list[dict[str, str]]:
    """A table sheet as dicts keyed by its header row; every value as text (dates ISO)."""
    rows = list(load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))
    header = [_text(c) for c in rows[0]]
    return [dict(zip(header, (_text(v) for v in row))) for row in rows[1:]]


def every_row(path: Path) -> list[dict[str, str]]:
    """All target-list rows: the main sheet, then the Ended sheet (the CSV holds the main sheet only)."""
    return read_sheet(path, "Target list") + read_sheet(path, "Ended")


def load_header(path: Path, sheet: str) -> list[str]:
    first = next(load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))
    return [_text(c) for c in first]


def sheet_text(path: Path, sheet: str) -> str:
    """Every non-empty cell of a sheet, one line per row, cells joined by ' | '."""
    lines = []
    for row in load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True):
        cells = [_text(v) for v in row if v not in (None, "")]
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)
