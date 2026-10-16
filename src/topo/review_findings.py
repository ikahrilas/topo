"""Review findings: the founder's spot-check marks, kept in the write-once curated file
data/curated/review_findings.csv.

    fain, state, recipient, program, mark, field, note, date, reviewer

A target-list row's finding carries its FAIN and State; a Missing-sheet finding carries
State, Recipient and Program and no FAIN. `mark` is ok / wrong / stale (or `missing`),
`field` names the column that is wrong or stale, `note` is at most `[spot_check]
note_max_chars` characters, `date` is the review date given to the ingest, and
`reviewer` is always `founder`. Every other column of the returned workbook is dropped,
so no extra personal data comes in.

The file is append-only: ingest opens it for appending and never rewrites a byte; a
finding already recorded is not appended again. Findings never change the overrides
file or the frozen answer key: each `wrong` Host mark is checked against the evidence
before a reviewer writes an override by hand.
"""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

import pandera.polars as pa
import polars as pl
from openpyxl import load_workbook

from topo.config import Config
from topo.curated import _read, curated_path
from topo.guards import StructuralBreak, check_schema
from topo.outputs.spot_check import FIELD, FIELDS_SHEET, LOOKS_RIGHT, MARKS, NOTE, SpotCheckSettings
from topo.outputs.target_list import ENDED_SHEET, LAST_FOUNDER_MARK, MAIN_SHEET, MISSING_COLUMNS, MISSING_SHEET

FINDINGS = "review_findings.csv"
REVIEWER = "founder"
MISSING = "missing"  # the mark of a Missing-sheet finding
COLUMNS = ("fain", "state", "recipient", "program", "mark", "field", "note", "date", "reviewer")


class ReviewFindings(pa.DataFrameModel):
    """One spot-check mark per row (glossary: Review finding)."""

    fain: str = pa.Field(nullable=True)
    state: str = pa.Field(nullable=True)
    recipient: str = pa.Field(nullable=True)
    program: str = pa.Field(nullable=True)
    mark: str = pa.Field(isin=[*MARKS, MISSING])
    field: str = pa.Field(nullable=True)
    note: str = pa.Field(nullable=True)
    date: str = pa.Field(str_matches=r"^\d{4}-\d{2}-\d{2}$")
    reviewer: str = pa.Field(isin=[REVIEWER])

    class Config:
        strict = True

    @pa.dataframe_check(error="a finding names a FAIN, or (Missing) a State with a Recipient or Program")
    def names_a_grant(cls, data: pa.PolarsData) -> pl.LazyFrame:
        return data.lazyframe.select(
            pl.col("fain").is_not_null() | (
                (pl.col("mark") == MISSING) & pl.col("state").is_not_null()
                & (pl.col("recipient").is_not_null() | pl.col("program").is_not_null())
            )
        )


def findings_path(config: Config) -> Path:
    return curated_path(config, FINDINGS)


def read_review_findings(config: Config) -> pl.DataFrame:
    """The findings file, checked against its schema; a missing file reads as empty."""
    return _read(findings_path(config), ReviewFindings, "Review findings")


def read_returned_workbook(path: Path, review_date: dt.date, settings: SpotCheckSettings) -> pl.DataFrame:
    """The findings in a returned spot-check copy: allowed columns only, reviewer `founder`.

    A mark outside ok / wrong / stale, a Field that is not a target-list column, or a Note
    longer than the limit stops the ingest.
    """
    if not path.exists():
        raise StructuralBreak(f"returned workbook not found: {path}")
    book = load_workbook(path, read_only=True, data_only=True)
    for sheet in (MAIN_SHEET, ENDED_SHEET, MISSING_SHEET, FIELDS_SHEET):
        if sheet not in book.sheetnames:
            raise StructuralBreak(f"{path.name} is not a spot-check copy: no {sheet!r} sheet")
    fields = {_text(row[0]) for row in book[FIELDS_SHEET].iter_rows(values_only=True)} - {None}
    date = review_date.isoformat()
    found: list[dict[str, str | None]] = []
    for sheet in (MAIN_SHEET, ENDED_SHEET):
        for row in _rows(book[sheet], (LOOKS_RIGHT, FIELD, NOTE, "FAIN", "State"), path):
            if not any(row[c] for c in (LOOKS_RIGHT, FIELD, NOTE)):
                continue
            where = f"{path.name}, {sheet}, FAIN {row['FAIN']}"
            mark = (row[LOOKS_RIGHT] or "").lower()
            if mark not in MARKS:
                raise StructuralBreak(f"{where}: \"{LOOKS_RIGHT}\" must be one of {', '.join(MARKS)}, not {mark!r}")
            if row[FIELD] is not None and row[FIELD] not in fields - {LAST_FOUNDER_MARK}:
                raise StructuralBreak(f"{where}: Field {row[FIELD]!r} is not a target-list column")
            found.append({
                "fain": row["FAIN"], "state": row["State"], "recipient": None, "program": None, "mark": mark,
                "field": row[FIELD], "note": _note(row[NOTE], settings, where), "date": date, "reviewer": REVIEWER,
            })
    for row in _rows(book[MISSING_SHEET], MISSING_COLUMNS, path):
        if not any(row.values()):
            continue
        where = f"{path.name}, {MISSING_SHEET}, {row['Recipient'] or row['Program']}"
        found.append({
            "fain": None, "state": row["State"], "recipient": row["Recipient"], "program": row["Program"],
            "mark": MISSING, "field": None, "note": _note(row["Note"], settings, where), "date": date,
            "reviewer": REVIEWER,
        })
    frame = pl.DataFrame(found, schema={c: pl.String for c in COLUMNS})
    return check_schema(ReviewFindings, frame, f"Review findings in {path.name}")


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _note(value: str | None, settings: SpotCheckSettings, where: str) -> str | None:
    if value is not None and len(value) > settings.note_max_chars:
        raise StructuralBreak(f"{where}: Note is longer than {settings.note_max_chars} characters")
    return value


def _rows(sheet, wanted, path: Path) -> list[dict[str, str | None]]:
    """The wanted columns of a sheet, as text; every other column is never read."""
    rows = sheet.iter_rows(values_only=True)
    header = [_text(c) for c in next(rows, ())]
    missing = [c for c in wanted if c not in header]
    if missing:
        raise StructuralBreak(f"{path.name}, {sheet.title}: lacks columns {', '.join(missing)}")
    index = {c: header.index(c) for c in wanted}
    return [
        {c: _text(row[i]) if i < len(row) else None for c, i in index.items()}
        for row in rows
    ]


def append_findings(config: Config, findings: pl.DataFrame) -> int:
    """Append findings not already recorded; returns the rows appended. Never rewrites the file."""
    existing = read_review_findings(config)
    seen = {tuple(r) for r in existing.select(COLUMNS).iter_rows()}
    new = [r for r in findings.select(COLUMNS).iter_rows() if r not in seen]
    new = list(dict.fromkeys(new))
    if not new:
        return 0
    path = findings_path(config)
    path.parent.mkdir(parents=True, exist_ok=True)
    fresh = not path.exists() or path.stat().st_size == 0
    if not fresh and not path.read_bytes().endswith(b"\n"):
        raise StructuralBreak(f"{path.name} does not end with a newline; fix it by hand before appending")
    with path.open("a", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        if fresh:
            writer.writerow(COLUMNS)
        writer.writerows(["" if v is None else v for v in row] for row in new)
    return len(new)


def ingest(config: Config, workbook: Path, review_date: dt.date) -> int:
    """Read a returned spot-check copy and append its findings; returns the rows appended."""
    settings = SpotCheckSettings.from_section(config.section("spot_check"))
    return append_findings(config, read_returned_workbook(workbook, review_date, settings))


def last_founder_marks(findings: pl.DataFrame) -> pl.DataFrame:
    """fain -> "Last founder mark": the latest-dated finding on each FAIN (file order breaks ties)."""
    marks = findings.filter(pl.col("fain").is_not_null()).with_row_index("order")
    shown = pl.concat_str(
        pl.col("mark"),
        pl.when(pl.col("field").is_not_null()).then(pl.lit(" (") + pl.col("field") + pl.lit(")")).otherwise(pl.lit("")),
        pl.lit(" on "),
        pl.col("date"),
        pl.when(pl.col("note").is_not_null()).then(pl.lit(": ") + pl.col("note")).otherwise(pl.lit("")),
    )
    return (
        marks.sort("date", "order")
        .group_by("fain", maintain_order=True)
        .agg(shown.last().alias(LAST_FOUNDER_MARK))
    )
