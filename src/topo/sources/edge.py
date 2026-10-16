"""NCES EDGE: postsecondary campus geocodes and composite school district polygons.

Raw layout: data/raw/edge/<Pull date>/<file names from the configured URLs>,
both zips as NCES ships them. The geocode zip holds an XLSX (or CSV) with
UNITID, NAME, STATE, LAT, LON and SCHOOLYEAR; the district zip holds one
shapefile whose features carry GEOID (= LEAID except pseudo-districts), NAME,
MTFCC (G5400 elementary, G5410 secondary, G5420 unified, G5430 administrative)
and SCHOOLYEAR. The shapefile is read with DuckDB spatial.

Cleaned tables:
    edge_postsecondary_geocodes  one row per UNITID with a campus point (lat/lon, NAD83)
                                 and its school year; rows with no point are dropped
                                 (`no_coordinates`).
    edge_school_districts        one row per district polygon: LEAID, name, district
                                 type, school year and the geometry as WKB. Land or water
                                 the state left without a district (SDLEA 99997) is
                                 dropped (`no_school_district_defined`).

Without any EDGE pull, `clean` is a no-op: no college Host then gets a Campus district.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import urllib.request
import zipfile
from collections.abc import Callable
from pathlib import Path
from urllib.parse import urlparse

import pandera.polars as pa
import polars as pl
from openpyxl import load_workbook

from topo import provenance, run_log, spatial
from topo.config import Config, Layout
from topo.guards import StructuralBreak, check_schema, check_unique
from topo.sources.usaspending import PulledFile

SOURCE = "edge"
GEOCODES_TABLE = "edge_postsecondary_geocodes"
DISTRICTS_TABLE = "edge_school_districts"
DROP_NO_COORDINATES = "no_coordinates"
DROP_NO_DISTRICT_DEFINED = "no_school_district_defined"
NO_DISTRICT_DEFINED_CODE = "99997"

ELEMENTARY, SECONDARY, UNIFIED, ADMINISTRATIVE = "elementary", "secondary", "unified", "administrative"
DISTRICT_TYPES = {"G5400": ELEMENTARY, "G5410": SECONDARY, "G5420": UNIFIED, "G5430": ADMINISTRATIVE}

LICENCE_CITATION = (
    "NCES Education Demographic and Geographic Estimates (EDGE), "
    "https://nces.ed.gov/programs/edge/; public domain (\"all information contained herein is in the public domain\")"
)

GEOCODE_COLUMNS = {"UNITID": "unitid", "NAME": "name", "STATE": "state", "LAT": "lat", "LON": "lon",
                   "SCHOOLYEAR": "school_year"}


class EdgePostsecondaryGeocodes(pa.DataFrameModel):
    """EDGE campus points: one row per UNITID."""

    unitid: str = pa.Field(unique=True)
    name: str = pa.Field(nullable=True)
    state: str = pa.Field(nullable=True)
    lat: float = pa.Field(ge=-90, le=90)
    lon: float = pa.Field(ge=-180, le=180)
    school_year: str = pa.Field(nullable=True)
    raw_file: str
    pull_date: dt.date

    class Config:
        strict = True


class EdgeSchoolDistricts(pa.DataFrameModel):
    """EDGE district polygons: one row per (LEAID, district type)."""

    leaid: str
    district_name: str = pa.Field(nullable=True)
    district_type: str = pa.Field(isin=list(DISTRICT_TYPES.values()))
    school_year: str = pa.Field(nullable=True)
    wkb: pl.Binary
    raw_file: str
    pull_date: dt.date

    class Config:
        strict = True
        unique = ["leaid", "district_type"]


def pull_dates(layout: Layout) -> list[dt.date]:
    root = layout.raw_source(SOURCE)
    if not root.is_dir():
        return []
    found = []
    for child in root.iterdir():
        try:
            found.append(dt.date.fromisoformat(child.name))
        except ValueError:
            continue
    return sorted(found)


def _file_name(url: str) -> str:
    return Path(urlparse(url).path).name


def _urls(config: Config) -> dict[str, str]:
    settings = config.section(SOURCE)
    return {"geocodes": settings["postsecondary_geocodes_url"], "districts": settings["school_districts_url"]}


Download = Callable[[str, Path], None]


def _download(url: str, dest: Path) -> None:
    with urllib.request.urlopen(url, timeout=1800) as response, dest.open("wb") as out:
        while chunk := response.read(1 << 20):
            out.write(chunk)


def pull(config: Config, pull_date: dt.date, *, download: Download = _download) -> list[PulledFile]:
    """Fetch the EDGE geocode and district zips (serially) into data/raw/edge/<Pull date>/."""
    folder = config.layout.raw_source(SOURCE) / pull_date.isoformat()
    folder.mkdir(parents=True, exist_ok=True)
    pulled = []
    for url in _urls(config).values():
        dest = folder / _file_name(url)
        download(url, dest)
        data = dest.read_bytes()
        pulled.append(
            PulledFile(
                source=SOURCE, request=f"GET {url}", url=url, path=dest, pull_date=pull_date,
                bytes=len(data), sha256=hashlib.sha256(data).hexdigest(), licence_citation=LICENCE_CITATION,
            )
        )
    return pulled


# ------------------------------------------------------------------ geocodes


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_geocode_rows(path: Path) -> pl.DataFrame:
    """The geocode table inside the zip (XLSX preferred, else CSV), all strings."""
    with zipfile.ZipFile(path) as z:
        members = z.namelist()
        xlsx = [m for m in members if m.lower().endswith(".xlsx")]
        csvs = [m for m in members if m.lower().endswith(".csv")]
        if xlsx:
            sheet = load_workbook(io.BytesIO(z.read(xlsx[0])), read_only=True).worksheets[0]
            rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            header, body = [h.upper() for h in rows[0]], rows[1:]
            raw = pl.DataFrame([dict(zip(header, r)) for r in body], schema={h: pl.String for h in header}, orient="row") \
                if body else pl.DataFrame(schema={h: pl.String for h in header})
        elif csvs:
            raw = pl.read_csv(z.read(csvs[0]), infer_schema=False)
            raw = raw.rename({c: c.upper() for c in raw.columns})
        else:
            raise StructuralBreak(f"{path.name} holds no XLSX or CSV geocode table")
    missing = sorted(set(GEOCODE_COLUMNS) - set(raw.columns))
    if missing:
        raise StructuralBreak(f"{path.name} lacks columns: {', '.join(missing)}")
    return raw.select(pl.col(src).replace("", None).alias(dst) for src, dst in GEOCODE_COLUMNS.items())


def clean_geocodes(raw: pl.DataFrame, *, raw_file: str, pull_date: dt.date) -> pl.DataFrame:
    points = (
        raw.with_columns(pl.col("lat").cast(pl.Float64, strict=False), pl.col("lon").cast(pl.Float64, strict=False))
        .filter(pl.col("lat").is_not_null() & pl.col("lon").is_not_null())
        .with_columns(pl.lit(raw_file).alias("raw_file"), pl.lit(pull_date).alias("pull_date"))
        .select(list(EdgePostsecondaryGeocodes.to_schema().columns))
        .sort("unitid")
    )
    check_unique(points, ["unitid"], "EDGE postsecondary geocodes")
    return check_schema(EdgePostsecondaryGeocodes, points, "EDGE postsecondary geocodes")


# ----------------------------------------------------------------- districts


def _shapefile_member(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        shp = [m for m in z.namelist() if m.lower().endswith(".shp")]
    if len(shp) != 1:
        raise StructuralBreak(f"{path.name} should hold one district shapefile; found {len(shp)}")
    return shp[0]


def _read_districts(path: Path, raw_file: str, pull_date: dt.date) -> tuple[int, int, pl.DataFrame]:
    """(rows in, rows with no district defined, cleaned polygons) from the zipped shapefile."""
    con = spatial.connect()
    source = spatial.sql_string(f"/vsizip/{path.resolve().as_posix()}/{_shapefile_member(path)}")
    columns = {c[0].upper() for c in con.execute(f"DESCRIBE SELECT * FROM ST_Read({source})").fetchall()}
    missing = sorted({"GEOID", "NAME", "MTFCC", "SCHOOLYEAR"} - columns)
    if missing:
        raise StructuralBreak(f"{raw_file} lacks columns: {', '.join(missing)}")
    rows = con.execute(
        f"SELECT GEOID, NAME, MTFCC, SCHOOLYEAR, ST_AsWKB(geom) FROM ST_Read({source})"
    ).fetchall()
    con.close()
    unknown = sorted({r[2] for r in rows} - set(DISTRICT_TYPES))
    if unknown:
        raise StructuralBreak(f"{raw_file} has district polygons of unknown MTFCC: {', '.join(map(str, unknown))}")
    undefined = [r for r in rows if str(r[0])[2:] == NO_DISTRICT_DEFINED_CODE]
    kept = [r for r in rows if str(r[0])[2:] != NO_DISTRICT_DEFINED_CODE]
    polygons = pl.DataFrame(
        {
            "leaid": [r[0] for r in kept],
            "district_name": [r[1] for r in kept],
            "district_type": [DISTRICT_TYPES[r[2]] for r in kept],
            "school_year": [r[3] for r in kept],
            "wkb": [bytes(r[4]) for r in kept],
        },
        schema={"leaid": pl.String, "district_name": pl.String, "district_type": pl.String,
                "school_year": pl.String, "wkb": pl.Binary},
    ).with_columns(pl.lit(raw_file).alias("raw_file"), pl.lit(pull_date).alias("pull_date")).sort(
        "leaid", "district_type"
    )
    check_unique(polygons, ["leaid", "district_type"], "EDGE school districts")
    return len(rows), len(undefined), check_schema(EdgeSchoolDistricts, polygons, "EDGE school districts")


# --------------------------------------------------------------------- clean


def clean(config: Config, pull_date: dt.date | None = None) -> pl.DataFrame:
    """Clean one EDGE pull (default: the latest). No pull at all: no-op, empty frame."""
    layout = config.layout
    if pull_date is None:
        dates = pull_dates(layout)
        if not dates:
            return _empty(EdgePostsecondaryGeocodes)
        pull_date = dates[-1]
    folder = layout.raw_source(SOURCE) / pull_date.isoformat()
    names = {key: _file_name(url) for key, url in _urls(config).items()}
    for name in names.values():
        if not (folder / name).exists():
            raise StructuralBreak(f"EDGE pull lacks {name} ({folder})")
    geocodes_path, districts_path = folder / names["geocodes"], folder / names["districts"]
    input_hash = provenance.record_raw_inputs(config, SOURCE, [geocodes_path, districts_path])

    raw_points = _read_geocode_rows(geocodes_path)
    points = clean_geocodes(raw_points, raw_file=names["geocodes"], pull_date=pull_date)
    no_point = raw_points.filter(
        pl.col("lat").cast(pl.Float64, strict=False).is_null() | pl.col("lon").cast(pl.Float64, strict=False).is_null()
    ).height
    run_log.record(
        config, layer="clean", source=GEOCODES_TABLE, rows_in=raw_points.height, rows_out=points.height,
        drops={DROP_NO_COORDINATES: no_point}, pull_date=pull_date, input_hash=input_hash,
    )

    rows_in, undefined, polygons = _read_districts(districts_path, names["districts"], pull_date)
    run_log.record(
        config, layer="clean", source=DISTRICTS_TABLE, rows_in=rows_in, rows_out=polygons.height,
        drops={DROP_NO_DISTRICT_DEFINED: undefined}, pull_date=pull_date, input_hash=input_hash,
    )
    layout.cleaned.mkdir(parents=True, exist_ok=True)
    points.write_parquet(layout.cleaned_table(GEOCODES_TABLE))
    polygons.write_parquet(layout.cleaned_table(DISTRICTS_TABLE))
    return points


def _empty(model: type[pa.DataFrameModel]) -> pl.DataFrame:
    return pl.DataFrame(schema={name: col.dtype.type for name, col in model.to_schema().columns.items()})


def read_geocodes(config: Config) -> pl.DataFrame:
    """Campus points; empty when EDGE was never pulled and cleaned."""
    path = config.layout.cleaned_table(GEOCODES_TABLE)
    return pl.read_parquet(path) if path.exists() else _empty(EdgePostsecondaryGeocodes)


def read_school_districts(config: Config) -> pl.DataFrame:
    """District polygons (WKB); empty when EDGE was never pulled and cleaned."""
    path = config.layout.cleaned_table(DISTRICTS_TABLE)
    return pl.read_parquet(path) if path.exists() else _empty(EdgeSchoolDistricts)
