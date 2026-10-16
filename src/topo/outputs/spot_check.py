"""The founder spot-check copy: a frozen copy of one pull's target-list workbook for marking,
delivered as output/spot_check_<Pull date>.xlsx and never overwritten.

Sheets: "Spot check" (instructions and the names warning), States (each state's Grant
count, to pick the 2–3 states to check), README (as the target list's), "Target list" and
Ended (the target-list rows with three editable mark columns in front: "Looks right?"
(ok / wrong / stale), Field (the column that is wrong or stale) and Note (at most
`note_max_chars` characters)), Missing (State, Recipient, Program, Note: Grants the founder
knows exist but doesn't see) and a hidden Fields sheet listing the column names Field may take.

Every cell except the mark columns and the Missing sheet is locked; "Last founder mark" is
read-only. `topo ingest-findings` reads the returned copy (topo.review_findings).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from topo.outputs.delivery import delivery_path, write_new
from topo.outputs.target_list import (
    ENDED_SHEET,
    MAIN_SHEET,
    MISSING_COLUMNS,
    MISSING_SHEET,
    README_SHEET,
    ReadmeFacts,
    _write_readme,
    split_sheets,
)

STEM = "spot_check"
INSTRUCTIONS_SHEET = "Spot check"
STATES_SHEET = "States"
FIELDS_SHEET = "Fields"
LOOKS_RIGHT = "Looks right?"
FIELD = "Field"
NOTE = "Note"
MARK_COLUMNS = (LOOKS_RIGHT, FIELD, NOTE)
MARKS = ("ok", "wrong", "stale")
NAMES_WARNING = (
    "Do not enter client or personal names. Name organizations, Programs and fields only; "
    "anything else is dropped when the marks are read back."
)
EDITABLE_FILL = PatternFill("solid", fgColor="FFF2CC")


@dataclass(frozen=True)
class SpotCheckSettings:
    """[spot_check] in config.toml. ASSUMPTION: the founder checks 2–3 states."""

    min_states: int = 2
    max_states: int = 3
    note_max_chars: int = 100

    @classmethod
    def from_section(cls, section: dict[str, Any]) -> "SpotCheckSettings":
        return cls(
            min_states=int(section.get("min_states", 2)),
            max_states=int(section.get("max_states", 3)),
            note_max_chars=int(section.get("note_max_chars", 100)),
        )


def state_counts(rows: pl.DataFrame) -> pl.DataFrame:
    """Grants per State over every target-list row (main and Ended sheets)."""
    return rows.group_by("State").agg(pl.len().alias("Grants")).sort("State", nulls_last=True)


def write_spot_check(
    rows: pl.DataFrame, readme: ReadmeFacts, settings: SpotCheckSettings, output_dir: Path
) -> Path:
    """Write the spot-check copy of the target-list rows to the next free delivery path."""
    path = delivery_path(output_dir, STEM, readme.pull_date, "xlsx")
    write_new(path, _workbook_bytes(rows, readme, settings))
    return path


def _workbook_bytes(rows: pl.DataFrame, readme: ReadmeFacts, settings: SpotCheckSettings) -> bytes:
    main, ended = split_sheets(rows)
    book = Workbook()
    _write_instructions(book.active, readme, settings)
    _write_states(book.create_sheet(STATES_SHEET), state_counts(rows))
    _write_readme(book.create_sheet(), readme)
    fields = book.create_sheet(FIELDS_SHEET)
    for name in rows.columns:
        fields.append([name])
    fields.sheet_state = "hidden"
    fields.protection.sheet = True
    for name, frame in ((MAIN_SHEET, main), (ENDED_SHEET, ended)):
        _write_marking_table(book.create_sheet(name), frame, len(rows.columns), settings)
    missing = book.create_sheet(MISSING_SHEET)
    missing.append(list(MISSING_COLUMNS))
    for cell in missing[1]:
        cell.font = Font(bold=True)
    # README and States stay read-only too; Missing is the founder's to fill.
    for sheet in (book[INSTRUCTIONS_SHEET], book[STATES_SHEET], book[README_SHEET]):
        sheet.protection.sheet = True
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_instructions(sheet, readme: ReadmeFacts, settings: SpotCheckSettings) -> None:
    sheet.title = INSTRUCTIONS_SHEET
    lines = [
        ("Founder spot check", ""),
        ("Warning", NAMES_WARNING),
        ("Pull date", readme.pull_date.isoformat()),
        ("As-of date", readme.as_of.isoformat()),
        ("1. Pick states", f"Pick {settings.min_states}–{settings.max_states} states you know well; the "
                           f"{STATES_SHEET} sheet gives each state's Grant count."),
        ("2. Mark rows", f"On the {MAIN_SHEET} and {ENDED_SHEET} sheets, fill \"{LOOKS_RIGHT}\" with "
                         f"{' / '.join(MARKS)} for the rows of those states. Leave other rows blank."),
        ("3. Name the field", f"For wrong or stale rows, pick in {FIELD} the column that is wrong or out of date."),
        ("4. Note", f"Optional, at most {settings.note_max_chars} characters: what you know instead."),
        ("5. Missing Grants", f"On the {MISSING_SHEET} sheet, add Grants you know exist but don't see "
                              "(State, Recipient, Program, Note)."),
        ("Last founder mark", "Your mark on that FAIN from an earlier spot check; read-only."),
        ("Afterwards", "Marks come back as Review findings. They never change a Host link or any other "
                       "field by themselves: each wrong mark is checked against the evidence first."),
    ]
    for term, text in lines:
        sheet.append([term, text or None])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["B2"].font = Font(bold=True, color="C00000")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 120


def _write_states(sheet, counts: pl.DataFrame) -> None:
    sheet.append(counts.columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in counts.iter_rows():
        sheet.append(list(row))


def _write_marking_table(sheet, frame: pl.DataFrame, field_count: int, settings: SpotCheckSettings) -> None:
    sheet.append([*MARK_COLUMNS, *frame.columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in frame.iter_rows():
        sheet.append([None, None, None, *row])
    last = sheet.max_row
    for col in range(1, len(MARK_COLUMNS) + 1):
        for row in range(2, last + 1):
            cell = sheet.cell(row=row, column=col)
            cell.protection = Protection(locked=False)
            cell.fill = EDITABLE_FILL
    if last >= 2:
        looks = DataValidation(type="list", formula1=f'"{",".join(MARKS)}"', allow_blank=True)
        field = DataValidation(type="list", formula1=f"{FIELDS_SHEET}!$A$1:$A${field_count}", allow_blank=True)
        note = DataValidation(
            type="textLength", operator="lessThanOrEqual", formula1=str(settings.note_max_chars), allow_blank=True,
            promptTitle="Note", prompt=f"At most {settings.note_max_chars} characters. No client or personal names.",
            showInputMessage=True,
        )
        for validation, col in ((looks, 1), (field, 2), (note, 3)):
            letter = get_column_letter(col)
            validation.add(f"{letter}2:{letter}{last}")
            sheet.add_data_validation(validation)
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    protection = sheet.protection
    protection.sheet = True
    protection.autoFilter = False  # filtering and sorting stay allowed on the locked sheet
    protection.sort = False
