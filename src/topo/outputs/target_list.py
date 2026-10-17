"""Target list: one row per Grant (FAIN) in the reporting window, as an XLSX
workbook (README, main, Ended and Missing sheets) plus a CSV of the main sheet.

Amount columns are USAspending obligations by action-date federal fiscal year,
labelled "FY2025", never by academic year; ED amounts never appear here.
Funding status, Latest readable status, Readable from, Lifecycle stage and the
2025–26 disruption marker come from the Funding status reader (#31). Host
institution, Host ID, Host confidence and Case label come from Host linking
(#32). The displayed label and Priority tier come from
topo.domain.priority_tier (#33). Re-competing soon, Predicted due date and
Evidence tier come from renewal timing (#40). Local context (#38): institution
type, enrollment, the Host's or Campus district's poverty rate, each with its
data year, and a reason code when null.

Rows sort by Priority tier, then Host confidence (matched first), then award
date, then State. Ended Grants go to the Ended sheet; every other row, whatever
its Host outcome, stays on the main sheet.
"""

from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import Workbook
from openpyxl.styles import Font

from topo.domain.cause_notes import cited
from topo.domain.funding_status import ENDED, NOT_YET_READABLE
from topo.domain.grant_universe import GrantUniverse
from topo.domain.local_context import CAMPUS_DISTRICT, HOST_DISTRICT, LocalContext
from topo.domain.host_linking import HOST_CONFIDENCE, MATCHED, NO_SINGLE_HOST, REVIEW, UNMATCHED
from topo.domain.priority_tier import RowFacts, label_row
from topo.domain.programs import PROGRAMS, is_dol_style
from topo.fiscal import ReportingWindow, fiscal_year, fy_end, fy_label
from topo.guards import StructuralBreak
from topo.outputs.delivery import delivery_paths, write_new

STEM = "target_list"
MAIN_SHEET = "Target list"
ENDED_SHEET = "Ended"
MISSING_SHEET = "Missing"
README_SHEET = "README"
MISSING_COLUMNS = ("State", "Recipient", "Program", "Note")
LAST_FOUNDER_MARK = "Last founder mark"  # the latest Review finding on the FAIN, read-only

FUNDING_COLUMNS = {  # joined grant_funding_status column -> display column
    "funding_status": "Funding status",
    "latest_readable_status": "Latest readable status",
    "readable_from": "Readable from",
    "lifecycle_stage": "Lifecycle stage",
    "disruption": "2025–26 disruption",
}
HOST_CONFIDENCE_ORDER = {HOST_CONFIDENCE[o]: i for i, o in enumerate((MATCHED, NO_SINGLE_HOST, REVIEW, UNMATCHED))}

CENSUS_NOTICE = "This product uses the Census Bureau Data API but is not endorsed or certified by the Census Bureau."
ATTRIBUTIONS = (
    "USAspending.gov, U.S. Department of Treasury, Bureau of the Fiscal Service, https://www.usaspending.gov. "
    "Accessed {pull_date}.",
    "Recipient names and addresses include Dun & Bradstreet (D&B) Open Data, republished from USAspending.gov "
    "with attribution to D&B (https://www.usaspending.gov/db-info).",
    "IPEDS, CCD, via Education Data Portal v. 0.27.0, Urban Institute, under ODC Attribution License "
    "(https://educationdata.urban.org; https://opendatacommons.org/licenses/by/1-0/).",
    "Campus points and school district boundaries: NCES EDGE (https://nces.ed.gov/programs/edge/), public domain. "
    "Child poverty: U.S. Census Bureau, Small Area Income and Poverty Estimates (SAIPE), school districts.",
)
GLOSSARY = (
    ("Grant", "One funded project for one project period, identified by its FAIN; a renewal is a new Grant."),
    ("Recipient", "The legal entity the award is made to, identified by its UEI."),
    ("Host institution", "The college (UNITID) or school district (LEAID) where the project runs; may differ "
                         "from the Recipient."),
    ("Host confidence", "matched; lower-confidence: no single Host (decided by rule to have no one Host); in review "
                        "(a person must choose); unmatched (no candidate found)."),
    ("Program", "The federal grant program, e.g. Upward Bound or GEAR UP state grants."),
    ("Budget period", "The stretch of a Grant, usually 12 months, that one funding action funds."),
    ("Funding status", "What happened to a Budget period's funding, read from funding actions: continued, funded "
                       "late, unfunded, reversed, or not yet readable. A pattern, never a cause."),
    ("Lifecycle stage", "Where a Grant sits in its project period: new, continuing, final budget period, ended; "
                        "re-competing soon when renewal timing flags it."),
    ("2025–26 disruption", "Funding missing, late or reversed for a Budget period starting 2025-05-01 to "
                           "2026-09-30, whether or not later restored; not proof of a policy cut."),
    ("Label", "One label per row, first match wins: Unfunded or Reversed, Re-competing soon, New, Final budget "
              "period, Ended, Continuing. \"not yet readable\" adds the date the latest Budget period becomes "
              "readable."),
    ("Priority tier", "Call-order bucket, not a score: 1 new within {months} months of the award; 2 re-competing "
                      "soon; 3 continuing; 4 unfunded or reversed (with its reason; not a call list)."),
    ("Predicted due date", "The expected application deadline of the Program's next competition: the prior "
                           "Cohort's published deadline plus 60 months (a prediction not yet tested). Evidence "
                           "tier: notice published or inferred. \"date unknown\" when no deadline is on file; "
                           "\"timing not predicted\" when the Program's cycle isn't 60 months, and for GEAR UP."),
    ("End date", "GEAR UP only: the Grant's expected end, when it falls within the lead time. An end date, "
                 "not a due date."),
    ("Local context", "Institution type, enrollment and the ages 5-17 poverty rate of the Host institution's "
                      "place, each with its data year; it describes where the Host sits, not who the project "
                      "serves. Never taken from the Recipient or the state. When blank, Context reason says why: "
                      "no_saipe_row, no_campus_district, no_enrollment_row, no_single_host, unmatched, in_review."),
    ("Campus district", "The school district a college Host's campus lies in; its poverty rate is labelled as the "
                        "campus's district, never as the served schools'. A campus in both an elementary and a "
                        "secondary or unified district shows the secondary or unified one and notes the overlap."),
    ("Cause note", "Shown only where a cited public source names the Grant's FAIN."),
    ("Last founder mark", "The founder's latest spot-check mark on the FAIN (ok, wrong or stale, with the field "
                          "and note). Read-only; a mark never changes any other field by itself."),
    ("Pull date", "The date the source data was fetched."),
    ("As-of date", "The date the data is read as of; Funding status is judged against it."),
)


@dataclass(frozen=True)
class AwaitedCohort:
    """A Program whose `fy` competition is known to have been awarded but whose new FAINs
    (`<FAIN prefix><yy>`, e.g. P044A26) may not be reported yet."""

    program: str
    fy: int

    @property
    def note(self) -> str:
        return f"{fy_label(self.fy)} new awards not yet reported"

    def fain_prefix(self) -> str:
        [prefix] = [p.fain_prefix for p in PROGRAMS if p.name == self.program]
        return f"{prefix}{self.fy % 100:02d}"


@dataclass(frozen=True)
class TargetListSettings:
    """[target_list] in config.toml."""

    new_award_months: int = 3
    provisional_lag_days: int = 90
    awaited_cohorts: tuple[AwaitedCohort, ...] = ()

    @classmethod
    def from_section(cls, section: dict[str, Any]) -> "TargetListSettings":
        return cls(
            new_award_months=int(section.get("new_award_months", 3)),
            provisional_lag_days=int(section.get("provisional_lag_days", 90)),
            awaited_cohorts=tuple(
                AwaitedCohort(str(c["program"]), int(c["fy"])) for c in section.get("awaited_cohorts", [])
            ),
        )


def provisional_fiscal_years(window: ReportingWindow, pull_date: dt.date, lag_days: int) -> list[int]:
    """Window FYs less than `lag_days` past their end on the Pull date (reporting lag)."""
    return [fy for fy in window.fiscal_years if (pull_date - fy_end(fy)).days < lag_days]


def provisional_note(fy: int, pull_date: dt.date) -> str:
    return f"{fy_label(fy)} provisional (as of {pull_date.isoformat()}; reporting lag)"


def awaited_cohort_notes(universe: GrantUniverse, settings: TargetListSettings) -> dict[tuple[str, int], str]:
    """(Program, FY) -> note, for awaited Cohorts none of whose FAINs has appeared yet.

    A Cohort has appeared once a `<FAIN prefix><yy>` FAIN is a Grant, or a DOL-style
    FAIN Grant of the Program starts in the Cohort's FY or later (DOL FAINs carry no ED prefix).
    """
    grants = universe.grants
    fains = grants["fain"].to_list()
    dol_starts = [
        (program, fiscal_year(start))
        for fain, agency, program, start in grants.select("fain", "awarding_agency_code", "program", "project_start").iter_rows()
        if is_dol_style(fain, agency)
    ]
    return {
        (c.program, c.fy): c.note
        for c in settings.awaited_cohorts
        if not any(f.startswith(c.fain_prefix()) for f in fains)
        and not any(program == c.program and fy >= c.fy for program, fy in dol_starts)
    }


def awaited_notes(universe: GrantUniverse, settings: TargetListSettings) -> dict[str, str]:
    """Program -> note, for awaited Cohorts none of whose FAINs has appeared yet."""
    return {program: note for (program, _), note in awaited_cohort_notes(universe, settings).items()}

TIMING_COLUMNS = {  # joined renewal_timing column -> display column (#40)
    "re_competing_soon_shown": "Re-competing soon",
    "predicted_due_date": "Predicted due date",
    "evidence_tier": "Evidence tier",
    "predicted_competition_fy": "Predicted competition FY",
    "timing": "Renewal timing",
    "end_date_fact": "End date",
}


POVERTY_BASIS = {CAMPUS_DISTRICT: "Campus district", HOST_DISTRICT: "Host district"}

LOCAL_CONTEXT_COLUMNS = {  # joined local_context column -> display column (#38; built in _local_context)
    "institution_type": "Institution type",
    "institution_type_year": "Institution type year",
    "enrollment": "Enrollment",
    "enrollment_year": "Enrollment year",
    "poverty_rate_5_17": "Poverty rate (ages 5-17)",
    "poverty_rate_of": "Poverty rate of",
    "poverty_rate_year": "Poverty rate year",
    "campus_district_year": "Campus district year",
    "campus_district_overlap": "Campus district overlap",
    "context_reason": "Context reason",
}


def _local_context(local_context: pl.DataFrame | None) -> pl.DataFrame:
    """Display columns of Local context; the poverty rate is labelled by whose district it is."""
    if local_context is None:
        local_context = pl.DataFrame(
            schema={name: col.dtype.type for name, col in LocalContext.to_schema().columns.items()}
        )
    return local_context.select(
        "fain",
        *(pl.col(src).alias(shown) for src, shown in LOCAL_CONTEXT_COLUMNS.items() if src != "poverty_rate_of"),
        pl.concat_str(
            pl.col("poverty_rate_basis").replace_strict(POVERTY_BASIS, default=None),
            pl.lit(": "), "poverty_district_name", pl.lit(" (LEAID "), "poverty_district_leaid", pl.lit(")"),
        ).alias("Poverty rate of"),
    )


def build_target_list(
    universe: GrantUniverse,
    as_of: dt.date,
    *,
    host_links: pl.DataFrame,
    funding: pl.DataFrame | None = None,
    funding_periods: pl.DataFrame | None = None,
    cause_notes: pl.DataFrame | None = None,
    pull_date: dt.date | None = None,
    settings: TargetListSettings = TargetListSettings(),
    timing: pl.DataFrame | None = None,
    local_context: pl.DataFrame | None = None,
    founder_marks: pl.DataFrame | None = None,
) -> pl.DataFrame:
    """Every target-list row (main and Ended sheets), sorted, in display column order.

    `host_links` is the joined Host-links table, one row per FAIN. `funding` is the
    joined grant_funding_status table (one row per FAIN) and `funding_periods` the
    joined funding_status table (one row per FAIN and Budget period). `cause_notes`
    is the curated cause-notes table (topo.domain.cause_notes). `timing` is the
    joined renewal_timing table (topo.domain.renewal_timing, one row per FAIN).
    `founder_marks` holds "Last founder mark" per FAIN (topo.review_findings.last_founder_marks),
    read-only. `local_context` is the joined local_context table (topo.domain.local_context.LocalContext).
    """
    if funding is None:
        funding = pl.DataFrame(schema={"fain": pl.String, **{c: pl.String for c in FUNDING_COLUMNS}})
    timing = _timing_columns(timing)
    grants = universe.grants
    if pull_date is None:
        pull_date = grants["pull_date"].max() if grants.height else as_of
    window = universe.window
    labels = [fy_label(fy) for fy in window.fiscal_years]
    amounts = universe.fy_obligations.with_columns(
        pl.col("fiscal_year").map_elements(fy_label, return_dtype=pl.String).alias("fy")
    )
    wide = (
        amounts.pivot(on="fy", index="fain", values="obligation")
        if amounts.height
        else pl.DataFrame(schema={"fain": pl.String})
    )
    for label in labels:
        if label not in wide.columns:
            wide = wide.with_columns(pl.lit(None, dtype=pl.Float64).alias(label))
    latest = (
        amounts.sort("fiscal_year")
        .group_by("fain")
        .agg(
            pl.col("fy").last().alias("Latest FY"),
            pl.col("fiscal_year").last().alias("latest_fiscal_year"),
            pl.col("obligation").last().alias("Latest FY obligation"),
            pl.col("obligation").sum().alias("Cumulative obligation"),
        )
    )
    # Every Grant keeps its row whatever its Host outcome (review and unmatched included).
    missing = grants.join(host_links, on="fain", how="anti")
    if missing.height:
        raise StructuralBreak(f"{missing.height} Grants have no Host link: {', '.join(missing['fain'].head(10))}")
    hosts = host_links.select(
        "fain",
        pl.col("host_name").alias("Host institution"),
        pl.concat_str("host_kind", pl.lit(" "), "host_id").alias("Host ID"),
        pl.col("host_confidence").alias("Host confidence"),
        pl.col("case_label").alias("Case label"),
    )
    notes = cited(cause_notes) if cause_notes is not None else pl.DataFrame(
        schema={"fain": pl.String, "Cause note": pl.String}
    )
    rows = (
        grants.join(hosts, on="fain", how="left")
        .join(_award_dates(universe.grant_actions), on="fain", how="left")
        .join(latest, on="fain", how="left")
        .join(wide.select("fain", *labels), on="fain", how="left")
        .join(funding.select("fain", *FUNDING_COLUMNS), on="fain", how="left")
        .join(_latest_readable_periods(funding_periods), on="fain", how="left")
        .join(notes, on="fain", how="left")
        .join(timing, on="fain", how="left")
        .join(_local_context(local_context), on="fain", how="left")
        .join(_founder_marks(founder_marks), on="fain", how="left")
        .with_columns(pl.col(labels).fill_null(0.0), pl.col("Cumulative obligation").fill_null(0.0))
    )
    rows = rows.with_columns(
        _labels(rows, as_of, settings),
        _data_notes(rows, provisional_fiscal_years(window, pull_date, settings.provisional_lag_days), pull_date,
                    awaited_notes(universe, settings)),
    ).unnest("row_label")
    return (
        rows.sort(
            pl.col("priority_tier"),
            pl.col("Host confidence").replace_strict(HOST_CONFIDENCE_ORDER, default=len(HOST_CONFIDENCE_ORDER)),
            pl.col("award_date"),
            pl.col("recipient_state"),
            pl.col("fain"),
            nulls_last=True,
        )
        .select(
            pl.col("priority_tier").alias("Priority tier"),
            pl.col("label").alias("Label"),
            pl.coalesce("timing_lifecycle_stage", "lifecycle_stage").alias("Lifecycle stage"),
            pl.col("funding_status").alias("Funding status"),
            pl.col("latest_readable_status").alias("Latest readable status"),
            pl.col("readable_from").alias("Readable from"),
            pl.col("disruption").alias("2025–26 disruption"),
            pl.col("tier_reason").alias("Tier reason"),
            *(pl.col(src).alias(shown) for src, shown in TIMING_COLUMNS.items()),
            "Cause note",
            pl.col("recipient_name").alias("Recipient"),
            pl.col("recipient_uei").alias("Recipient UEI"),
            "Host institution",
            "Host ID",
            pl.col("recipient_state").alias("State"),
            pl.col("program").alias("Program"),
            pl.col("fain").alias("FAIN"),
            "Host confidence",
            "Case label",
            pl.col("project_start").alias("Project start"),
            pl.col("project_end").alias("Project end"),
            pl.col("award_date").alias("Award date"),
            "Latest FY",
            "Latest FY obligation",
            "Cumulative obligation",
            *labels,
            *LOCAL_CONTEXT_COLUMNS.values(),
            pl.col("data_note").alias("Data note"),
            LAST_FOUNDER_MARK,
            pl.col("pull_date").alias("Pull date"),
            pl.lit(as_of).alias("As-of date"),
        )
    )


def _timing_columns(timing: pl.DataFrame | None) -> pl.DataFrame:
    """The renewal-timing fields per FAIN, ready to show (Predicted competition FY as "FY2027")."""
    if timing is None:
        timing = pl.DataFrame(schema={"fain": pl.String, "re_competing_soon": pl.Boolean,
                                      "predicted_due_date": pl.Date, "evidence_tier": pl.String,
                                      "predicted_competition_fy": pl.Int64, "timing": pl.String,
                                      "end_date_fact": pl.String, "lifecycle_stage": pl.String})
    return timing.select(
        "fain",
        pl.col("lifecycle_stage").alias("timing_lifecycle_stage"),  # renewal timing adds "re-competing soon"
        pl.col("re_competing_soon"),
        pl.when(pl.col("re_competing_soon")).then(pl.lit("yes")).otherwise(pl.lit("no"))
        .alias("re_competing_soon_shown"),
        "predicted_due_date",
        "evidence_tier",
        (pl.lit("FY") + pl.col("predicted_competition_fy").cast(pl.String)).alias("predicted_competition_fy"),
        "timing",
        "end_date_fact",
    )
def _founder_marks(marks: pl.DataFrame | None) -> pl.DataFrame:
    if marks is None:
        return pl.DataFrame(schema={"fain": pl.String, LAST_FOUNDER_MARK: pl.String})
    return marks.select("fain", LAST_FOUNDER_MARK)


def _award_dates(grant_actions: pl.DataFrame) -> pl.DataFrame:
    """The award date of each Grant: its earliest A1 action, else its earliest action."""
    return grant_actions.group_by("fain").agg(
        pl.coalesce(
            pl.col("action_date").filter(pl.col("action_type") == "A1").min(), pl.col("action_date").min()
        ).alias("award_date")
    )


def _latest_readable_periods(periods: pl.DataFrame | None) -> pl.DataFrame:
    schema = {"fain": pl.String, "readable_period_start": pl.Date, "readable_period_end": pl.Date}
    if periods is None or not periods.height:
        return pl.DataFrame(schema=schema)
    return (
        periods.filter(pl.col("funding_status") != NOT_YET_READABLE)
        .sort("budget_period_start")
        .group_by("fain")
        .agg(
            pl.col("budget_period_start").last().alias("readable_period_start"),
            pl.col("budget_period_end").last().alias("readable_period_end"),
        )
    )


def _labels(rows: pl.DataFrame, as_of: dt.date, settings: TargetListSettings) -> pl.Series:
    values = []
    for r in rows.iter_rows(named=True):
        start, end = r["readable_period_start"], r["readable_period_end"]
        row = label_row(
            RowFacts(
                funding_status=r["funding_status"],
                readable_from=r["readable_from"],
                latest_readable_status=r["latest_readable_status"],
                latest_readable_period=(start, end) if start is not None else None,
                lifecycle_stage=r["lifecycle_stage"],
                award_date=r["award_date"],
                re_competing_soon=bool(r["re_competing_soon"]),
            ),
            as_of,
            settings.new_award_months,
        )
        values.append({"label": row.label, "priority_tier": row.priority_tier, "tier_reason": row.tier_reason})
    return pl.Series(
        "row_label", values,
        dtype=pl.Struct({"label": pl.String, "priority_tier": pl.Int64, "tier_reason": pl.String}),
    )


def _data_notes(
    rows: pl.DataFrame, provisional: list[int], pull_date: dt.date, awaited: dict[str, str]
) -> pl.Series:
    """Per row: the provisional label when its Latest FY is provisional, and the awaited-Cohort note."""
    values = []
    for fy, program in rows.select("latest_fiscal_year", "program").iter_rows():
        notes = []
        if fy in provisional:
            notes.append(provisional_note(fy, pull_date))
        if program in awaited:
            notes.append(awaited[program])
        values.append("; ".join(notes) or None)
    return pl.Series("data_note", values, dtype=pl.String)


def split_sheets(rows: pl.DataFrame) -> tuple[pl.DataFrame, pl.DataFrame]:
    """(main sheet, Ended sheet). Only ended Grants leave the main sheet."""
    ended = pl.col("Lifecycle stage") == ENDED
    return rows.filter(~ended.fill_null(False)), rows.filter(ended.fill_null(False))


@dataclass(frozen=True)
class ReadmeFacts:
    pull_date: dt.date
    as_of: dt.date
    window: ReportingWindow
    settings: TargetListSettings
    provisional: list[int] = field(default_factory=list)
    awaited: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class TargetListFiles:
    xlsx: Path
    csv: Path


def write_target_list(rows: pl.DataFrame, readme: ReadmeFacts, output_dir: Path) -> TargetListFiles:
    """Write the workbook and the CSV of its main sheet, under one free version suffix."""
    main, ended = split_sheets(rows)
    paths = delivery_paths(output_dir, STEM, readme.pull_date, ("xlsx", "csv"))
    write_new(paths["xlsx"], _workbook_bytes(main, ended, readme))
    write_new(paths["csv"], main.write_csv().encode())
    return TargetListFiles(xlsx=paths["xlsx"], csv=paths["csv"])


def _workbook_bytes(main: pl.DataFrame, ended: pl.DataFrame, readme: ReadmeFacts) -> bytes:
    book = Workbook()
    _write_readme(book.active, readme)
    _write_table(book.create_sheet(MAIN_SHEET), main)
    _write_table(book.create_sheet(ENDED_SHEET), ended)
    missing = book.create_sheet(MISSING_SHEET)
    missing.append(list(MISSING_COLUMNS))
    for cell in missing[1]:
        cell.font = Font(bold=True)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_table(sheet, frame: pl.DataFrame) -> None:
    sheet.append(frame.columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in frame.iter_rows():
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _write_readme(sheet, facts: ReadmeFacts) -> None:
    sheet.title = README_SHEET
    window = f"{fy_label(facts.window.start_fy)}–{fy_label(facts.window.end_fy)}"
    lines: list[tuple[str, str]] = [
        ("TRIO and GEAR UP target list", ""),
        ("Pull date", facts.pull_date.isoformat()),
        ("As-of date", facts.as_of.isoformat()),
        ("Reporting window", f"{window} (federal fiscal years; amounts are USAspending obligations)"),
        ("Sheets", f"{MAIN_SHEET}: live Grants, sorted by Priority tier, lower Host confidence below matched rows "
                   f"within a tier. {ENDED_SHEET}: ended Grants. {MISSING_SHEET}: Grants you know exist but don't "
                   "see (State, Recipient, Program, Note)."),
    ]
    lines += [("Provisional", provisional_note(fy, facts.pull_date)) for fy in facts.provisional]
    lines += [("Not yet reported", f"{program}: {note}") for program, note in sorted(facts.awaited.items())]
    lines.append(("", ""))
    lines.append(("Glossary", ""))
    lines += [(term, text.format(months=facts.settings.new_award_months)) for term, text in GLOSSARY]
    lines.append(("", ""))
    lines.append(("Census notice", CENSUS_NOTICE))
    lines += [("Attribution", text.format(pull_date=facts.pull_date.isoformat())) for text in ATTRIBUTIONS]
    for term, text in lines:
        sheet.append([term or None, text or None])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 120
